"""DOCX Exporter — Phase 4.

Renders a ReportContent to a .docx file using python-docx.
Charts embedded as PNG images generated via matplotlib.
All data comes from the passed-in ReportContent object.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Optional

from app.services.report.assembler import ReportContent

logger = logging.getLogger(__name__)


def _build_chart_bytes(series_list, title: str) -> Optional[bytes]:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        fig, ax = plt.subplots(figsize=(7, 3.5))
        has_data = False
        for s in series_list:
            xs = [dp.period_label for dp in s.data_points if dp.value is not None]
            ys = [dp.value for dp in s.data_points if dp.value is not None]
            if xs and ys:
                ax.plot(xs, ys, marker="o", label=s.entity_name)
                has_data = True
        if not has_data:
            plt.close(fig)
            return None
        ax.set_title(title, fontsize=10)
        ax.set_xlabel("Period")
        ax.set_ylabel(series_list[0].unit if series_list else "Value")
        ax.legend(fontsize=7)
        plt.xticks(rotation=30, fontsize=7)
        plt.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=120)
        plt.close(fig)
        buf.seek(0)
        return buf.read()
    except Exception as exc:
        logger.warning("Chart generation failed: %s", exc)
        return None


async def export_docx(content: ReportContent, output_path: Path) -> None:
    """Render ReportContent to a .docx file at output_path."""
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX export.") from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = DocxDocument()

    # Title
    title = doc.add_heading("Mining Intelligence Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Scope: {content.scope.label}")
    doc.add_paragraph(f"Period: {content.scope.period_start} to {content.scope.period_end}")
    doc.add_paragraph(f"Generated: {content.generated_at.strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_paragraph(f"Model: {content.model_used}")
    doc.add_page_break()

    # Executive Summary
    _add_narrative_section(doc, content.executive_summary.heading, content.executive_summary.text)

    # Production Overview Table
    doc.add_heading(content.production_overview.heading, level=2)
    if content.production_overview.rows:
        cols = ["Entity", "Metric", "Period", "Value", "Unit", "Citation"]
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Light Shading Accent 1"
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(cols):
            hdr_cells[i].text = col
        for row in content.production_overview.rows:
            cells = table.add_row().cells
            cells[0].text = str(row.get("entity", ""))
            cells[1].text = str(row.get("metric", ""))
            cells[2].text = str(row.get("period", ""))
            cells[3].text = str(row.get("value", ""))
            cells[4].text = str(row.get("unit", ""))
            cells[5].text = str(row.get("citation_key", ""))

    # Historical Trends chart
    doc.add_heading(content.historical_trends.heading, level=2)
    if content.historical_trends.series:
        chart_bytes = _build_chart_bytes(
            content.historical_trends.series,
            title=f"Trends — {content.scope.label}",
        )
        if chart_bytes:
            img_stream = io.BytesIO(chart_bytes)
            doc.add_picture(img_stream, width=Inches(5.5))

    # Comparative Analysis
    doc.add_heading(content.comparative_analysis.heading, level=2)
    if content.comparative_analysis.rows:
        cols = ["Entity", "Metric", "Period", "Value", "Unit"]
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Light Shading Accent 1"
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(cols):
            hdr_cells[i].text = col
        for row in content.comparative_analysis.rows:
            cells = table.add_row().cells
            for i, key in enumerate(("entity", "metric", "period", "value", "unit")):
                cells[i].text = str(row.get(key, ""))

    # Data Quality Warnings
    doc.add_heading(content.data_quality_warnings.heading, level=2)
    if content.data_quality_warnings.open_flags:
        doc.add_heading(f"Open Flags ({len(content.data_quality_warnings.open_flags)})", level=3)
        for f in content.data_quality_warnings.open_flags[:20]:
            doc.add_paragraph(
                f"[{f.get('severity','?').upper()}] {f.get('flag_type','')} — "
                f"{f.get('entity_name','')} / {f.get('metric','')}",
                style="List Bullet",
            )
    if content.data_quality_warnings.open_conflicts:
        doc.add_heading(f"Open Conflicts ({len(content.data_quality_warnings.open_conflicts)})", level=3)
        for c in content.data_quality_warnings.open_conflicts[:20]:
            doc.add_paragraph(
                f"Conflict {str(c.get('conflict_id',''))[:8]}... — "
                f"{c.get('entity_name','')} / {c.get('metric','')} — "
                f"delta {c.get('delta_pct','?')}%",
                style="List Bullet",
            )

    # Recommendations
    if content.recommendations:
        _add_narrative_section(doc, content.recommendations.heading, content.recommendations.text)

    # Citations
    doc.add_heading("Source Citations", level=2)
    for cit in content.citations:
        doc.add_paragraph(
            f"[{cit.claim_key}] {cit.document_filename} (p.{cit.page_number or '?'}) — {cit.excerpt[:150]}",
            style="List Bullet",
        )

    doc.save(str(output_path))
    logger.info("DOCX exported: %s", output_path)


def _add_narrative_section(doc, heading: str, text: str) -> None:
    doc.add_heading(heading, level=2)
    if text:
        for para in text.split("\n\n"):
            para = para.strip()
            if para:
                doc.add_paragraph(para)
    else:
        doc.add_paragraph("(Narrative unavailable — see structured data sections.)")
