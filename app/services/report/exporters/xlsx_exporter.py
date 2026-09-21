"""XLSX Exporter — Phase 4.

Renders a ReportContent to a multi-worksheet .xlsx file using openpyxl.
One worksheet per major section. Charts embedded as matplotlib PNG images.
All data comes from the passed-in ReportContent object.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Optional

from app.services.report.assembler import ReportContent

logger = logging.getLogger(__name__)


async def export_xlsx(content: ReportContent, output_path: Path) -> None:
    """Render ReportContent to a .xlsx file at output_path."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as XLImage
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for XLSX export.") from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1A3A5C")
    alt_fill    = PatternFill(fill_type="solid", fgColor="EEF2F7")
    bold        = Font(bold=True)
    center      = Alignment(horizontal="center")

    # ------------------------------------------------------------------ #
    # Sheet 1 — Metadata
    # ------------------------------------------------------------------ #
    ws_meta = wb.active
    ws_meta.title = "Report Info"
    _write_row(ws_meta, 1, ["Field", "Value"], header_font, header_fill)
    rows_meta = [
        ["Report ID", str(content.report_id)],
        ["Generated At", content.generated_at.isoformat()],
        ["Scope Label", content.scope.label],
        ["Period Start", str(content.scope.period_start)],
        ["Period End", str(content.scope.period_end)],
        ["Model Used", content.model_used],
        ["DQ Warning Count", content.dq_warning_count],
        ["Citation Count", len(content.citations)],
        ["Generation Time (s)", content.generation_time_seconds],
    ]
    for i, row in enumerate(rows_meta, start=2):
        ws_meta.cell(i, 1, row[0])
        ws_meta.cell(i, 2, row[1])
    ws_meta.column_dimensions["A"].width = 25
    ws_meta.column_dimensions["B"].width = 40

    # ------------------------------------------------------------------ #
    # Sheet 2 — Production Overview
    # ------------------------------------------------------------------ #
    ws_prod = wb.create_sheet("Production Overview")
    cols = ["Entity", "Metric", "Period", "Value", "Unit", "Citation"]
    _write_row(ws_prod, 1, cols, header_font, header_fill)
    for i, row in enumerate(content.production_overview.rows, start=2):
        fill = alt_fill if i % 2 == 0 else None
        _write_row(ws_prod, i, [
            row.get("entity", ""),
            row.get("metric", ""),
            row.get("period", ""),
            row.get("value", ""),
            row.get("unit", ""),
            row.get("citation_key", ""),
        ], fill=fill)
    _auto_width(ws_prod, cols)

    # ------------------------------------------------------------------ #
    # Sheet 3 — Historical Trends
    # ------------------------------------------------------------------ #
    ws_trends = wb.create_sheet("Historical Trends")
    cols_t = ["Entity", "Metric", "Period Label", "Period Start", "Period End", "Value", "Unit", "Has Conflict", "Citation"]
    _write_row(ws_trends, 1, cols_t, header_font, header_fill)
    r = 2
    for s in content.historical_trends.series:
        for dp in s.data_points:
            fill = alt_fill if r % 2 == 0 else None
            _write_row(ws_trends, r, [
                s.entity_name, s.metric,
                dp.period_label,
                str(dp.period_start) if dp.period_start else "",
                str(dp.period_end)   if dp.period_end   else "",
                dp.value,
                s.unit or "",
                "Yes" if dp.has_conflict else "No",
                dp.citation_key or "",
            ], fill=fill)
            r += 1

    # Embed chart image
    chart_bytes = _build_chart_bytes(content.historical_trends.series, f"Trends — {content.scope.label}")
    if chart_bytes:
        img = XLImage(io.BytesIO(chart_bytes))
        img.anchor = f"A{r + 2}"
        ws_trends.add_image(img)
    _auto_width(ws_trends, cols_t)

    # ------------------------------------------------------------------ #
    # Sheet 4 — Comparative Analysis
    # ------------------------------------------------------------------ #
    ws_comp = wb.create_sheet("Comparative Analysis")
    cols_c = ["Entity", "Metric", "Period", "Value", "Unit"]
    _write_row(ws_comp, 1, cols_c, header_font, header_fill)
    for i, row in enumerate(content.comparative_analysis.rows, start=2):
        fill = alt_fill if i % 2 == 0 else None
        _write_row(ws_comp, i, [
            row.get("entity", ""), row.get("metric", ""),
            row.get("period", ""), row.get("value", ""), row.get("unit", ""),
        ], fill=fill)
    _auto_width(ws_comp, cols_c)

    # ------------------------------------------------------------------ #
    # Sheet 5 — Data Quality Warnings
    # ------------------------------------------------------------------ #
    ws_dq = wb.create_sheet("DQ Warnings")
    ws_dq.cell(1, 1, "Open Validation Flags").font = bold
    cols_f = ["Flag Type", "Severity", "Entity", "Metric", "Detail"]
    _write_row(ws_dq, 2, cols_f, header_font, header_fill)
    for i, f in enumerate(content.data_quality_warnings.open_flags, start=3):
        _write_row(ws_dq, i, [
            f.get("flag_type", ""), f.get("severity", ""),
            f.get("entity_name", ""), f.get("metric", ""),
            str(f.get("detail", "")),
        ])
    offset = max(4, len(content.data_quality_warnings.open_flags) + 4)
    ws_dq.cell(offset, 1, "Open Conflicts").font = bold
    cols_co = ["Conflict ID", "Entity", "Metric", "Period", "Delta %", "Fact A", "Fact B"]
    _write_row(ws_dq, offset + 1, cols_co, header_font, header_fill)
    for i, c in enumerate(content.data_quality_warnings.open_conflicts, start=offset + 2):
        _write_row(ws_dq, i, [
            str(c.get("conflict_id", ""))[:8] + "...",
            c.get("entity_name", ""), c.get("metric", ""), c.get("period", ""),
            c.get("delta_pct", ""),
            str(c.get("fact_a_id", ""))[:8] + "...",
            str(c.get("fact_b_id", ""))[:8] + "...",
        ])

    # ------------------------------------------------------------------ #
    # Sheet 6 — Citations
    # ------------------------------------------------------------------ #
    ws_cit = wb.create_sheet("Citations")
    cols_ci = ["Claim Key", "Document", "Page", "Excerpt", "Confidence", "Fact ID"]
    _write_row(ws_cit, 1, cols_ci, header_font, header_fill)
    for i, cit in enumerate(content.citations, start=2):
        fill = alt_fill if i % 2 == 0 else None
        _write_row(ws_cit, i, [
            cit.claim_key,
            cit.document_filename,
            cit.page_number or "",
            cit.excerpt[:200],
            cit.extraction_confidence or "",
            str(cit.normalized_fact_id),
        ], fill=fill)
    _auto_width(ws_cit, cols_ci)

    wb.save(str(output_path))
    logger.info("XLSX exported: %s", output_path)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _write_row(ws, row_idx: int, values, font=None, fill=None) -> None:
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row_idx, col_idx, val)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill


def _auto_width(ws, cols) -> None:
    for i, col_name in enumerate(cols, start=1):
        letter = __import__("openpyxl.utils", fromlist=["get_column_letter"]).get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, len(col_name) + 4)


def _build_chart_bytes(series_list, title: str) -> Optional[bytes]:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        fig, ax = plt.subplots(figsize=(7, 3.5))
        has_data = False
        for s in series_list:
            xs = [dp.period_label for dp in s.data_points if dp.value is not None]
            ys = [dp.value for dp in s.data_points if dp.value is not None]
            if xs and ys:
                ax.plot(xs, ys, marker="o", label=s.entity_name)
                has_data = True
        if not has_data:
            plt.close(fig)
            return None
        ax.set_title(title, fontsize=10)
        ax.set_xlabel("Period")
        ax.legend(fontsize=7)
        plt.xticks(rotation=30, fontsize=7)
        plt.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=120)
        plt.close(fig)
        buf.seek(0)
        return buf.read()
    except Exception as exc:
        logger.warning("Chart generation failed: %s", exc)
        return None
