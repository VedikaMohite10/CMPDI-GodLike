"""Extractor for XLSX and CSV files.

Both formats are natively structured — no OCR is applied.
Each sheet (XLSX) or file (CSV) becomes one Table with a single logical page.

XLSX: One logical page per sheet; sheet name is recorded in position metadata.
CSV : Single page, single table.
"""
import asyncio
import csv
import io
import logging
from typing import List

import openpyxl

from app.services.extraction.base import (
    AbstractExtractor,
    ExtractionResult,
    Table,
    TextBlock,
)

logger = logging.getLogger(__name__)


def _extract_xlsx_sync(file_bytes: bytes) -> ExtractionResult:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    tables: List[Table] = []

    for sheet_idx, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        rows_data: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            str_row = [str(cell) if cell is not None else "" for cell in row]
            rows_data.append(str_row)

        if not rows_data:
            continue

        # Skip entirely-empty sheets
        if all(all(c == "" for c in r) for r in rows_data):
            continue

        headers = rows_data[0]
        rows = rows_data[1:]
        raw_structure = {
            "headers": headers,
            "rows": rows,
            "shape": [len(rows), len(headers)],
        }
        tables.append(
            Table(
                page_number=sheet_idx + 1,      # one page per sheet
                table_index=0,
                raw_structure=raw_structure,
                caption=sheet_name,
                position={"sheet": sheet_name},
            )
        )

    wb.close()
    page_count = max((t.page_number for t in tables), default=1)
    return ExtractionResult(
        page_count=page_count,
        text_blocks=[],
        tables=tables,
        images=[],
        page_dimensions={i: (0.0, 0.0) for i in range(1, page_count + 1)},
    )


def _extract_csv_sync(file_bytes: bytes) -> ExtractionResult:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1")

    reader = csv.reader(io.StringIO(text))
    rows_data: List[List[str]] = list(reader)

    if not rows_data:
        return ExtractionResult(page_count=1, page_dimensions={1: (0.0, 0.0)})

    headers = rows_data[0]
    rows = rows_data[1:]
    raw_structure = {"headers": headers, "rows": rows, "shape": [len(rows), len(headers)]}

    table = Table(
        page_number=1,
        table_index=0,
        raw_structure=raw_structure,
    )
    return ExtractionResult(
        page_count=1,
        text_blocks=[],
        tables=[table],
        images=[],
        page_dimensions={1: (0.0, 0.0)},
    )


class SpreadsheetExtractor(AbstractExtractor):
    """Extractor for XLSX + CSV — no OCR, direct structured parse."""

    def __init__(self, file_type: str):
        self._file_type = file_type  # 'xlsx' | 'csv'

    async def extract(self, file_bytes: bytes) -> ExtractionResult:
        if self._file_type == "csv":
            return await asyncio.to_thread(_extract_csv_sync, file_bytes)
        return await asyncio.to_thread(_extract_xlsx_sync, file_bytes)
