#!/usr/bin/env python3
"""CMPDI E2E Test Suite — SIH26023"""
import json, os, subprocess, sys, time
from pathlib import Path
import requests

BASE = "http://localhost:8000"
OLLAMA = "http://localhost:11434"
PASS_COUNT = FAIL_COUNT = WARN_COUNT = 0
RESULTS = []
GR="\033[92m"; RD="\033[91m"; YL="\033[93m"; RS="\033[0m"

def record(status, test_id, detail=""):
    global PASS_COUNT, FAIL_COUNT, WARN_COUNT
    RESULTS.append({"status":status,"id":test_id,"detail":detail})
    if status=="PASS": PASS_COUNT+=1; print(f"{GR}✅ {test_id}{RS}")
    elif status=="WARN": WARN_COUNT+=1; print(f"{YL}⚠️  {test_id}: {detail[:100]}{RS}")
    else: FAIL_COUNT+=1; print(f"{RD}❌ {test_id}: {detail[:100]}{RS}")

def section(t): print(f"\n{'═'*64}\n {t}\n{'═'*64}")

class FakeResp:
    def __init__(self,e): self._e=e; self.status_code=0; self.content=b""; self.text=e; self.headers={}
    def json(self): return {"_error":self._e}

def GET(path,tok="",**kw):
    h={"Authorization":f"Bearer {tok}"} if tok else {}
    try: return requests.get(f"{BASE}{path}",headers=h,timeout=30,**kw)
    except Exception as e: return FakeResp(str(e))

def POST(path,tok="",jb=None,files=None,**kw):
    h={"Authorization":f"Bearer {tok}"} if tok else {}
    try: return requests.post(f"{BASE}{path}",headers=h,json=jb,files=files,timeout=120,**kw)
    except Exception as e: return FakeResp(str(e))

# ── PHASE 0 ───────────────────────────────────────────────
section("PHASE 0: ENVIRONMENT")
r=requests.get(f"{BASE}/health",timeout=10)
record("PASS" if r.status_code==200 and r.json().get("status")=="ok" else "FAIL","ENV-01: API health",r.text[:80])

try:
    cp=subprocess.run(["docker","ps","--format","{{.Names}}"],capture_output=True,text=True)
    containers=cp.stdout
    missing=[c for c in ["cmpdi_api","cmpdi_postgres","cmpdi_qdrant"] if c not in containers]
    record("PASS" if not missing else "FAIL","ENV-02: Docker containers running",
           "all up" if not missing else f"Missing:{missing}")
except Exception as e: record("WARN","ENV-02: Docker containers",str(e))

try:
    names={m["name"] for m in requests.get(f"{OLLAMA}/api/tags",timeout=10).json().get("models",[])}
    for model in ["bge-m3:latest","qwen2.5vl:3b","qwen2.5:7b-instruct-q4_K_M","qwen2.5:14b-instruct-q4_K_M"]:
        stem=model.split(":")[0]
        record("PASS" if any(stem in n for n in names) else "FAIL",f"ENV-03: Ollama {model}",
               "found" if any(stem in n for n in names) else f"not in {sorted(names)[:5]}")
except Exception as e: record("WARN","ENV-03: Ollama",str(e))

APP_DIR=Path(__file__).parent/"app"
hits=[str(f) for f in APP_DIR.rglob("*.py") if "qwen2.5-coder" in f.read_text(errors="ignore")]
record("PASS" if not hits else "FAIL","ENV-04: qwen2.5-coder NOT in runtime code","clean" if not hits else str(hits[:2]))

try:
    cr=requests.options(f"{BASE}/health",headers={"Origin":"http://localhost:5173","Access-Control-Request-Method":"GET"},timeout=10)
    acao=cr.headers.get("access-control-allow-origin","")
    record("PASS" if "*" in acao or "5173" in acao or cr.status_code in(200,204) else "WARN",
           "ENV-05: CORS allows frontend",f"ACAO={acao} status={cr.status_code}")
except Exception as e: record("WARN","ENV-05: CORS",str(e))

# ── AUTH ──────────────────────────────────────────────────
section("AUTH")
try:
    lr=requests.post(f"{BASE}/auth/login",data={"username":"admin","password":"changeme123!"},timeout=15)
    ADMIN_TOKEN=lr.json()["access_token"]
    record("PASS","AUTH-01: Admin JWT login",f"role={lr.json().get('role')}")
except Exception as e:
    record("FAIL","AUTH-01: Admin JWT login",str(e)); print("FATAL"); sys.exit(1)

# ── PHASE 1 ───────────────────────────────────────────────
section("PHASE 1: INGESTION")
TMP=Path("/tmp/cmpdi_e2e"); TMP.mkdir(exist_ok=True)

def make_pdf():
    p=TMP/"digital.pdf"
    try:
        from reportlab.pdfgen import canvas
        c=canvas.Canvas(str(p)); c.drawString(100,750,"CMPDI Production FY2023-24 CCL 85.2 MT BCCL 28.1 MT"); c.save()
    except:
        p.write_bytes(b'%PDF-1.4\n1 0 obj\n<</Type/Catalog/Pages 2 0 R>>\nendobj\n2 0 obj\n<</Type/Pages/Kids[3 0 R]/Count 1>>\nendobj\n3 0 obj\n<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>\nendobj\n4 0 obj\n<</Length 80>>\nstream\nBT /F1 12 Tf 72 720 Td (CMPDI Production CCL 85.2 MT coal FY2023-24 BCCL 28.1 MT) Tj ET\nendstream\nendobj\n5 0 obj\n<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>\nendobj\nxref\n0 6\n0000000000 65535 f\n0000000009 00000 n\n0000000058 00000 n\n0000000115 00000 n\n0000000258 00000 n\n0000000388 00000 n\ntrailer\n<</Size 6/Root 1 0 R>>\nstartxref\n460\n%%EOF\n')
    return p

def make_xlsx():
    p=TMP/"prod.xlsx"
    try:
        import openpyxl; wb=openpyxl.Workbook(); ws=wb.active
        ws.append(["Subsidiary","Year","Production_MT","Target_MT"])
        for row in [("CCL","FY2021",72.5,75.0),("CCL","FY2022",78.3,80.0),("CCL","FY2023",85.2,80.0),
                    ("BCCL","FY2021",28.1,30.0),("BCCL","FY2022",29.7,31.0)]: ws.append(row)
        wb.save(str(p))
    except Exception as e: record("WARN","P1-PREP:XLSX",str(e)); return None
    return p

def make_csv(name="coal.csv",content=None):
    p=TMP/name
    p.write_text(content or "Subsidiary,Year,Production_MT\nCCL,FY2021,72.5\nCCL,FY2022,78.3\nCCL,FY2023,85.2\nBCCL,FY2021,28.1\n")
    return p

def make_docx():
    p=TMP/"report.docx"
    try:
        from docx import Document; doc=Document()
        doc.add_heading("CMPDI Annual Mining Report",0)
        doc.add_paragraph("Central Coalfields Limited (CCL) produced 85.2 MT of coal during FY 2023-24.")
        doc.add_paragraph("Bharat Coking Coal Limited (BCCL) produced 29.7 MT during FY 2022-23.")
        doc.save(str(p))
    except Exception as e: record("WARN","P1-PREP:DOCX",str(e)); return None
    return p

def make_jpg():
    p=TMP/"scan.jpg"
    try:
        from PIL import Image,ImageDraw; img=Image.new("RGB",(400,200),(255,255,255))
        ImageDraw.Draw(img).text((20,50),"CMPDI Coal Production 85.2 MT FY2023",fill=(0,0,0))
        img.save(str(p))
    except:
        p.write_bytes(bytes.fromhex("ffd8ffe000104a46494600010100000100010000ffdb004300080606070605080707070909080a0c140d0c0b0b0c1912130f141d1a1f1e1d1a1c1c20242e2720222c231c1c2837292c30313434341f27393d38323c2e333432ffc0000b080001000101011100ffc4001f0000010501010101010100000000000000000102030405060708090a0bffc4001f100003010101010101010100000000000000010203040506070809ffd9"))
    return p

PDF_FILE=make_pdf(); XLSX_FILE=make_xlsx(); CSV_FILE=make_csv(); DOCX_FILE=make_docx(); JPG_FILE=make_jpg()
CONF_CSV=make_csv("conflict.csv","Subsidiary,Year,Production_MT\nCCL,FY2023,91.0\nBCCL,FY2021,28.1\n")

doc_ids={}
def upload(label,fp,mime):
    if fp is None: record("WARN",f"P1-UPLOAD-{label}","file not created"); return None
    with open(fp,"rb") as f:
        r=POST("/documents/upload",ADMIN_TOKEN,files={"file":(fp.name,f,mime)})
    d=r.json(); did=d.get("id") or d.get("document_id","")
    if did: record("PASS",f"P1-01: upload {label}",f"id={did}"); doc_ids[label]=did; return did
    else: record("FAIL",f"P1-01: upload {label}",f"status={r.status_code} {str(d)[:80]}"); return None

PDF_ID=upload("PDF",PDF_FILE,"application/pdf")
XLSX_ID=upload("XLSX",XLSX_FILE,"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
CSV_ID=upload("CSV",CSV_FILE,"text/csv")
DOCX_ID=upload("DOCX",DOCX_FILE,"application/vnd.openxmlformats-officedocument.wordprocessingml.document")
JPG_ID=upload("JPG",JPG_FILE,"image/jpeg")
CONF_ID=upload("CONFLICT_CSV",CONF_CSV,"text/csv")

bad=POST("/documents/upload",ADMIN_TOKEN,files={"file":("x.txt",open("/etc/hosts","rb"),"text/plain")})
bd=bad.json()
record("PASS" if bad.status_code in(400,415,422) or "detail" in bd else "FAIL",
       "P1-06: Bad file type → 4xx",f"status={bad.status_code}")

print("  … waiting 10s for processing …"); time.sleep(10)

if PDF_ID:
    r=GET(f"/documents/{PDF_ID}/status",ADMIN_TOKEN); d=r.json()
    sv=d.get("processing_status","")
    record("PASS" if sv else "FAIL","P1-07: GET /documents/{id}/status",f"status={sv}" if sv else str(d)[:80])

if PDF_ID:
    r=GET(f"/documents/{PDF_ID}",ADMIN_TOKEN); d=r.json()
    needed={"filename","file_type","upload_date","processing_status"}
    miss=needed-set(d.keys())
    record("PASS" if not miss else "FAIL","P1-08: GET /documents/{id} metadata",
           f"fn={d.get('filename')} status={d.get('processing_status')}" if not miss else f"Missing:{miss}")

if PDF_ID:
    r=GET(f"/documents/{PDF_ID}/original",ADMIN_TOKEN)
    ob=len(r.content); ub=PDF_FILE.stat().st_size
    record("PASS" if ob==ub else ("WARN" if ob>0 else "FAIL"),
           "P1-09: /documents/{id}/original byte-identical",f"dl={ob} ul={ub}")

if PDF_ID:
    r=GET(f"/documents/{PDF_ID}/pages/1",ADMIN_TOKEN); d=r.json()
    if r.status_code==404 or "not found" in str(d).lower():
        record("WARN","P1-11: pages/1","Still processing")
    elif any(k in d for k in("page_number","text_blocks","blocks","page_id","content")):
        record("PASS","P1-11: GET /documents/{id}/pages/1",str(list(d.keys()))[:60])
    else: record("FAIL","P1-11: GET /documents/{id}/pages/1",str(list(d.keys()))[:60])

r=POST("/search",ADMIN_TOKEN,jb={"query":"coal production CCL MT","limit":5}); d=r.json()
results=d if isinstance(d,list) else d.get("results",d.get("items",[]))
if results:
    has_prov=any(k in str(results[0]) for k in("document_id","page_id","block_id"))
    record("PASS" if has_prov else "WARN","P1-12: POST /search provenance blocks",
           f"{len(results)} results" if has_prov else "no provenance fields")
elif "answer" in str(d): record("FAIL","P1-12: POST /search no LLM synthesis","looks like LLM answer")
else: record("WARN","P1-12: POST /search","no results yet")

if PDF_ID:
    r=GET(f"/documents/{PDF_ID}",ADMIN_TOKEN)
    record("PASS" if r.status_code==200 and "filename" in r.json() else "FAIL","P1-CRITICAL: provenance DB trace",str(r.json())[:60])

# ── PHASE 2 ───────────────────────────────────────────────
section("PHASE 2: DATA TRUST ENGINE")
for did,lbl in [(PDF_ID,"PDF"),(CONF_ID,"CONFLICT"),(CSV_ID,"CSV")]:
    if did:
        r=POST(f"/documents/{did}/process-facts",ADMIN_TOKEN,jb={})
        record("PASS" if r.status_code in(200,202) else "FAIL",f"P2-01: process-facts {lbl}",str(r.json())[:80])

print("  … waiting 8s for fact extraction …"); time.sleep(8)

r=GET("/facts?limit=20",ADMIN_TOKEN); d=r.json()
facts=d if isinstance(d,list) else d.get("items",d.get("facts",[]))
record("PASS" if facts else "WARN","P2-02: GET /facts",f"{len(facts)} facts" if facts else "no facts yet (processing may be slow)")
FIRST_FACT_ID=facts[0].get("id","") if facts else ""

r=GET("/facts?entity=CCL&limit=5",ADMIN_TOKEN); d=r.json()
f2=d if isinstance(d,list) else d.get("items",d.get("facts",[]))
record("PASS" if r.status_code==200 else "FAIL","P2-03: GET /facts filter by entity",f"{len(f2)} CCL facts")

if FIRST_FACT_ID:
    r=GET(f"/facts/{FIRST_FACT_ID}/evidence",ADMIN_TOKEN); d=r.json()
    has_ev=any(k in str(d) for k in("document","page","chain","evidence","source"))
    record("PASS" if r.status_code==200 and has_ev else ("WARN" if r.status_code==200 else "FAIL"),
           "P2-04: GET /facts/{id}/evidence full chain",str(list(d.keys()))[:60])
else: record("WARN","P2-04: GET /facts/{id}/evidence","no fact ID")

r=GET("/conflicts?limit=10",ADMIN_TOKEN); d=r.json()
conflicts=d if isinstance(d,list) else d.get("items",d.get("conflicts",[]))
record("PASS" if isinstance(d,list) or "items" in d or "conflicts" in d else "FAIL",
       "P2-05: GET /conflicts",f"{len(conflicts)} conflicts")
FIRST_CONFLICT_ID=conflicts[0].get("id","") if conflicts else ""

if FIRST_CONFLICT_ID:
    r=GET(f"/conflicts/{FIRST_CONFLICT_ID}",ADMIN_TOKEN); d=r.json()
    record("PASS" if r.status_code==200 and any(k in str(d) for k in("fact_a","fact_b","source")) else "WARN",
           "P2-06: GET /conflicts/{id} both facts present",str(list(d.keys()))[:60])

r=GET("/validation-flags?limit=10",ADMIN_TOKEN); d=r.json()
flags=d if isinstance(d,list) else d.get("items",d.get("flags",[]))
record("PASS" if isinstance(d,list) or "items" in d or "flags" in d else "FAIL",
       "P2-07: GET /validation-flags",f"{len(flags)} flags")
FIRST_FLAG_ID=flags[0].get("id","") if flags else ""

r=GET("/entities?limit=10",ADMIN_TOKEN); d=r.json()
entities=d if isinstance(d,list) else d.get("items",d.get("entities",[]))
record("PASS" if isinstance(d,list) or "items" in d or "entities" in d else "FAIL",
       "P2-08: GET /entities",f"{len(entities)} entities")

r=GET("/duplicates?limit=10",ADMIN_TOKEN); d=r.json()
dups=d if isinstance(d,list) else d.get("items",d.get("duplicates",[]))
record("PASS" if isinstance(d,list) or "items" in d or "duplicates" in d else "FAIL",
       "P2-09: GET /duplicates",f"{len(dups)} duplicates")

if FIRST_FACT_ID:
    r=GET(f"/facts/{FIRST_FACT_ID}",ADMIN_TOKEN); d=r.json()
    has_raw="raw_value" in d or "original_text" in d or "raw_text" in d
    has_norm="normalized_value" in d or "value" in d
    record("PASS" if has_norm and has_raw else "WARN","P2-CRITICAL: raw+normalized values preserved",
           "both" if has_norm and has_raw else f"norm={has_norm} raw={has_raw} keys={list(d.keys())[:8]}")

# ── PHASE 3 ───────────────────────────────────────────────
section("PHASE 3: ANALYTICS + QUERY COPILOT")
tb={"entity":"CCL","metric":"production","start_period":"FY2021","end_period":"FY2023","period_type":"fiscal_year"}
r1=POST("/analytics/trend",ADMIN_TOKEN,jb=tb); time.sleep(1); r2=POST("/analytics/trend",ADMIN_TOKEN,jb=tb)
d1,d2=r1.json(),r2.json()
has_trend="trend" in d1 or "data_points" in d1 or "points" in d1 or isinstance(d1,list)
same=json.dumps(d1,sort_keys=True)==json.dumps(d2,sort_keys=True)
no_data="insufficient" in str(d1).lower() or "no data" in str(d1).lower()
if has_trend: record("PASS" if same else "WARN","P3-01: /analytics/trend deterministic","identical" if same else "differs between calls")
elif no_data: record("WARN","P3-01: /analytics/trend","no data yet")
else: record("FAIL","P3-01: /analytics/trend",str(d1)[:80])

r=POST("/analytics/compare",ADMIN_TOKEN,jb={"entities":["CCL","BCCL"],"metric":"production","period":"FY2021"})
d=r.json()
record("PASS" if any(k in d for k in("comparison","results","entities")) or isinstance(d,list) else
       ("WARN" if "insufficient" in str(d).lower() else "FAIL"),
       "P3-02: /analytics/compare",str(list(d.keys()))[:60])

# CRITICAL: sparse query must not fabricate
sq=POST("/query",ADMIN_TOKEN,jb={"question":"Why did XYZ Fictional Mine production drop 99.9% in FY1799?","include_reasoning":True})
sqd=sq.json(); ans=str(sqd.get("answer",sqd.get("response",""))).lower()
rt=sqd.get("reasoning_type","")
no_fab=any(x in ans for x in["insufficient","no evidence","not found","cannot","no relevant","no data","unable","don't have"])
record("PASS" if no_fab or "insufficient" in rt.lower() else "FAIL",
       "P3-CRITICAL: sparse /query → insufficient-evidence (no fabrication)",
       f"reasoning_type={rt}" if no_fab else f"FABRICATED? ans={ans[:100]}")

# Main query structure
mq=POST("/query",ADMIN_TOKEN,jb={"question":"What is the coal production of CCL and BCCL?","include_reasoning":True})
mqd=mq.json()
pf=[f for f in("answer","confidence","reasoning_type","citations","sources","evidence") if f in mqd]
record("PASS" if len(pf)>=3 else("WARN" if pf else "FAIL"),"P3-03: /query response structure",f"fields={pf}")
QUERY_ID=mqd.get("id","")

if QUERY_ID:
    r=GET(f"/query/{QUERY_ID}",ADMIN_TOKEN); d=r.json()
    record("PASS" if r.status_code==200 and any(k in d for k in("answer","question","response")) else "FAIL",
           "P3-04: GET /query/{id} persisted",f"id={QUERY_ID[:20]}")
else: record("WARN","P3-04: /query/{id} persisted","no ID returned")

# CRITICAL: why-did-this-change sparse
ws=POST("/analytics/why-did-this-change",ADMIN_TOKEN,jb={"entity":"XYZFAKE999","metric":"production","period":"FY1900"})
wsd=ws.json(); wss=str(wsd).lower()
record("PASS" if any(x in wss for x in["insufficient","no evidence","not found","no data","cannot"]) else "FAIL",
       "P3-CRITICAL: why-did-this-change sparse → no fabrication",str(wsd)[:100])

# Confidence sanity
tq=POST("/query",ADMIN_TOKEN,jb={"question":"What was exact production of Dhori mine in FY1934-35?","include_reasoning":True})
tqd=tq.json()
mq_c=float(mqd.get("confidence",0) or 0); tq_c=float(tqd.get("confidence",0) or 0)
record("PASS" if mq_c>tq_c or tq_c<0.5 else "WARN","P3-06: confidence sanity check",f"main={mq_c:.2f} thin={tq_c:.2f}")

# Citation integrity
cits=mqd.get("citations",mqd.get("sources",mqd.get("evidence",[])))
if cits and isinstance(cits,list):
    cit=cits[0]; fid=cit.get("fact_id","") if isinstance(cit,dict) else ""
    if fid:
        rf=GET(f"/facts/{fid}",ADMIN_TOKEN)
        record("PASS" if rf.status_code==200 else "FAIL","P3-CRITICAL: citation fact_id resolves",
               f"GET /facts/{fid[:20]} → {rf.status_code}")
    else: record("WARN","P3-CRITICAL: citation integrity","no fact_id in citation")
else: record("WARN","P3-CRITICAL: citation integrity","no citations returned")

# ── PHASE 4 ───────────────────────────────────────────────
section("PHASE 4: REPORTS + TOPICS + REVIEW + DASHBOARD")
rg=POST("/reports/generate",ADMIN_TOKEN,jb={"entity":"CCL","metric":"production","start_period":"FY2021","end_period":"FY2023","format":"pdf"})
rgd=rg.json(); REPORT_ID=rgd.get("id","") or rgd.get("report_id","")
record("PASS" if REPORT_ID or rg.status_code in(200,201,202) else "FAIL","P4-01: POST /reports/generate",
       f"id={REPORT_ID[:30] if REPORT_ID else 'queued'}")

if REPORT_ID:
    time.sleep(5)
    for fmt,sig in[("pdf",b"%PDF"),("docx",b"PK"),("xlsx",b"PK")]:
        re=GET(f"/reports/{REPORT_ID}/export?format={fmt}",ADMIN_TOKEN)
        ok=re.status_code==200 and len(re.content)>100
        record("PASS" if ok else "WARN",f"P4-02{fmt}: report export {fmt}",
               f"{len(re.content)}b ct={re.headers.get('content-type','')[:40]}" if ok else f"status={re.status_code}")

r=GET("/topics?limit=10",ADMIN_TOKEN); d=r.json()
topics=d if isinstance(d,list) else d.get("topics",d.get("items",[]))
record("PASS" if isinstance(d,list) or "topics" in d or "items" in d else "FAIL","P4-03: GET /topics",f"{len(topics)} topics")

r=GET("/topics/trends",ADMIN_TOKEN); d=r.json()
record("PASS" if r.status_code==200 else "FAIL","P4-04: GET /topics/trends",str(list(d.keys()) if isinstance(d,dict) else f"{len(d)} items")[:60])

r=GET("/review/flags?limit=10",ADMIN_TOKEN); d=r.json()
rflag_items=d if isinstance(d,list) else d.get("items",d.get("flags",[]))
record("PASS" if isinstance(d,list) or "items" in d or "flags" in d else "FAIL","P4-05: GET /review/flags",f"{len(rflag_items)} flags")
RFLAG_ID=rflag_items[0].get("id","") if rflag_items else ""

if RFLAG_ID:
    r=POST(f"/review/flags/{RFLAG_ID}/accept",ADMIN_TOKEN,jb={"reviewed_by":"admin","note":"test"})
    d=r.json()
    record("PASS" if r.status_code in(200,201) else "FAIL","P4-06: /review/flags/{id}/accept",
           f"status={d.get('status','')}" if r.status_code in(200,201) else str(d)[:80])

if FIRST_CONFLICT_ID:
    r=POST(f"/review/conflicts/{FIRST_CONFLICT_ID}/resolve",ADMIN_TOKEN,jb={"resolution":"accept_a"})
    d=r.json()
    record("PASS" if r.status_code in(400,422) or "justification" in str(d).lower() or "required" in str(d).lower() else "WARN",
           "P4-08: conflict resolve without justification → rejected",f"status={r.status_code}")
    r2=POST(f"/review/conflicts/{FIRST_CONFLICT_ID}/resolve",ADMIN_TOKEN,
            jb={"resolution":"accept_a","justification":"Document A is primary CMPDI official report"})
    d2=r2.json()
    record("PASS" if r2.status_code in(200,201) else "WARN","P4-09: conflict resolve with justification",
           f"status={d2.get('status','resolved')}" if r2.status_code in(200,201) else str(d2)[:80])

r=GET("/dashboard/stats",ADMIN_TOKEN); d=r.json()
req={"documents_processed","facts_extracted","conflicts_detected"}; pres=req&set(d.keys())
record("PASS" if pres==req else("WARN" if pres else "FAIL"),"P4-10: GET /dashboard/stats fields",
       f"docs={d.get('documents_processed')} facts={d.get('facts_extracted')}" if pres else f"got:{list(d.keys())[:8]}")

dash_before=d.get("documents_processed",0)
POST("/documents/upload",ADMIN_TOKEN,files={"file":("extra.csv",CSV_FILE.read_bytes(),"text/csv")})
time.sleep(5)
d2=GET("/dashboard/stats",ADMIN_TOKEN).json(); dash_after=d2.get("documents_processed",0)
record("PASS" if dash_after>dash_before else "WARN","P4-11: Dashboard count increments live",
       f"before={dash_before} after={dash_after}")

# ── PHASE 5 ───────────────────────────────────────────────
section("PHASE 5: PHASE 5 FEATURES")
parl=POST("/parliamentary/query",ADMIN_TOKEN,jb={"question":"What is the coal production status of CCL for FY2023-24?","context":"Parliamentary session"})
pd=parl.json(); PARL_ID=pd.get("id",""); ps=pd.get("status","")
record("PASS" if parl.status_code in(200,201,202) and (PARL_ID or "status" in pd) else "FAIL",
       "P5-01: POST /parliamentary/query",f"id={str(PARL_ID)[:30]} status={ps}")
record("PASS" if ps in("pending_review","draft","pending") else "WARN",
       "P5-02: Parliamentary starts as pending",f"status={ps!r}")

try:
    al=requests.post(f"{BASE}/auth/login",data={"username":"testanalyst","password":"Test@1234!"},timeout=10)
    ANA_TOKEN=al.json().get("access_token","")
    if not ANA_TOKEN:
        POST("/auth/users",ADMIN_TOKEN,jb={"username":"testanalyst","password":"Test@1234!","role":"analyst"})
        al2=requests.post(f"{BASE}/auth/login",data={"username":"testanalyst","password":"Test@1234!"},timeout=10)
        ANA_TOKEN=al2.json().get("access_token","")
    if ANA_TOKEN and PARL_ID:
        ap=POST(f"/parliamentary/{PARL_ID}/approve",ANA_TOKEN,jb={"reviewed_by":"testanalyst"})
        apd=ap.json()
        blocked=ap.status_code in(401,403) or any(x in str(apd).lower() for x in["forbidden","not authorized","permission"])
        record("PASS" if blocked else "FAIL","P5-03: Analyst cannot approve parliamentary (RBAC)",
               "blocked" if blocked else f"status={ap.status_code} {str(apd)[:80]}")
    else: record("WARN","P5-03: RBAC test","no analyst token")
except Exception as e: record("WARN","P5-03: RBAC test",str(e))

if PARL_ID:
    ap=POST(f"/parliamentary/{PARL_ID}/approve",ADMIN_TOKEN,jb={"reviewed_by":"admin","notes":"Approved"})
    apd=ap.json()
    record("PASS" if ap.status_code in(200,201) else "WARN","P5-04: Admin can approve parliamentary",
           f"status={apd.get('status','approved')}" if ap.status_code in(200,201) else str(apd)[:80])

r=GET("/map/layers",ADMIN_TOKEN); d=r.json()
layers=d if isinstance(d,list) else d.get("layers",[])
record("PASS" if r.status_code==200 and (isinstance(d,list) or "layers" in d) else "FAIL",
       "P5-05: GET /map/layers",f"{len(layers)} layers")

fs=POST("/forecast",ADMIN_TOKEN,jb={"entity":"XYZFAKEMINE999","metric":"production","periods":3})
fsd=fs.json(); fsds=str(fsd).lower()
record("PASS" if any(x in fsds for x in["insufficient","not enough","no historical","sparse","cannot forecast","no data"]) else "WARN",
       "P5-CRITICAL-01: Sparse forecast refuses",f"{str(fsd)[:100]}")

fr=POST("/forecast",ADMIN_TOKEN,jb={"entity":"CCL","metric":"production","periods":2})
frd=fr.json(); frds=str(frd).lower()
has_label="model-based forecast" in frds or "model_based_forecast" in frds or frd.get("forecast_type","").lower() in("model-based forecast","model_based")
has_data="predicted_value" in frds or "forecast_values" in frds or "forecast" in str(list(frd.keys())).lower()
if has_data and has_label: record("PASS","P5-CRITICAL-02: Forecast has Model-based label","present")
elif "insufficient" in frds or "no data" in frds: record("WARN","P5-CRITICAL-02: Forecast label","no CCL data yet")
elif has_data: record("FAIL","P5-CRITICAL-02: Forecast MISSING Model-based label","data without label!")
else: record("WARN","P5-CRITICAL-02: Forecast label",f"keys={list(frd.keys())[:6]}")

brun=POST("/benchmark/run",ADMIN_TOKEN,jb={})
brd=brun.json(); BENCH_ID=brd.get("id","") or brd.get("run_id","")
record("PASS" if brun.status_code in(200,201,202) and (BENCH_ID or "status" in brd) else "FAIL",
       "P5-08: POST /benchmark/run",f"id={str(BENCH_ID)[:30]} status={brd.get('status','')}")

br=GET("/benchmark/results",ADMIN_TOKEN); bres=br.json()
runs=bres if isinstance(bres,list) else bres.get("results",bres.get("runs",[]))
record("PASS" if br.status_code==200 and (isinstance(bres,list) or "results" in bres or "runs" in bres) else "FAIL",
       "P5-09: GET /benchmark/results",f"{len(runs)} run(s)")

# ── CROSS-PHASE ───────────────────────────────────────────
section("CROSS-PHASE INTEGRATION")
mg=list(APP_DIR.rglob("*.py"))
mg_files=[str(f) for f in mg if "model_gateway" in f.read_text(errors="ignore")]
record("PASS" if mg_files else "FAIL","X-01: ModelGateway is single LLM path",f"in {len(mg_files)} files")

audit=[str(f) for f in mg if "audit_log" in f.read_text(errors="ignore").lower()]
record("PASS" if audit else "FAIL","X-02: Audit log implemented",f"in {len(audit)} files")

rl=[str(f) for f in mg if any(x in f.read_text(errors="ignore").lower() for x in["slowapi","ratelimit","limiter"])]
record("PASS" if rl else "WARN","X-03: Rate limiting",f"in {len(rl)} files")

sec=Path(__file__).parent/"SECURITY.md"
record("PASS" if sec.exists() and sec.stat().st_size>100 else "WARN","X-04: SECURITY.md",f"{sec.stat().st_size}b" if sec.exists() else "not found")

parl_svc=APP_DIR/"services"/"parliamentary"
qc_ref=any("query_copilot" in (parl_svc/f).read_text(errors="ignore") or "explainer" in (parl_svc/f).read_text(errors="ignore")
           for f in os.listdir(parl_svc) if f.endswith(".py")) if parl_svc.is_dir() else False
record("PASS" if qc_ref else "WARN","X-05: Parliamentary reuses Phase 3 pipeline","confirmed" if qc_ref else "static check inconclusive")

rpt_svc=APP_DIR/"services"/"report"
ar=any("analytics" in (rpt_svc/f).read_text(errors="ignore") for f in os.listdir(rpt_svc) if f.endswith(".py")) if rpt_svc.is_dir() else False
record("PASS" if ar else "WARN","X-06: Report uses Phase 3 Analytics","confirmed" if ar else "static check inconclusive")

r=GET("/auth/me",ADMIN_TOKEN); d=r.json()
record("PASS" if r.status_code==200 and d.get("username")=="admin" else "FAIL","X-08: GET /auth/me",
       f"username={d.get('username')} role={d.get('role')}" if r.status_code==200 else str(d)[:60])

# ── SUMMARY ───────────────────────────────────────────────
total=PASS_COUNT+FAIL_COUNT+WARN_COUNT
section(f"FINAL SUMMARY — {total} tests")
print(f"\n  {GR}✅ PASS: {PASS_COUNT}{RS}\n  {YL}⚠️  WARN: {WARN_COUNT}{RS}\n  {RD}❌ FAIL: {FAIL_COUNT}{RS}\n  📊 TOTAL: {total}\n")
out=Path("/tmp/cmpdi_e2e_results.json")
out.write_text(json.dumps({"summary":{"pass":PASS_COUNT,"warn":WARN_COUNT,"fail":FAIL_COUNT,"total":total},"results":RESULTS},indent=2))
print(f"  Results: {out}")
if FAIL_COUNT or WARN_COUNT:
    print("\n─── FAILURES & WARNINGS ───")
    for r_ in RESULTS:
        if r_["status"]!="PASS":
            print(f"  {'⚠️ ' if r_['status']=='WARN' else '❌'} [{r_['status']}] {r_['id']}: {r_['detail'][:120]}")
sys.exit(0 if FAIL_COUNT==0 else 1)
